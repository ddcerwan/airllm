#!/usr/bin/env python3
"""Render the detailed AirLLM security report to Word (.docx) and PDF from report_content.BLOCKS.

Run:  python scripts/gen_detailed_report.py [base_output_path_without_extension]
Requires: python-docx (Word) and reportlab (PDF).
"""

import sys

from report_content import BLOCKS

# Reuse the Word styling helpers already written for the first report.
from gen_security_report import ACCENT, GREY, shade, style_header_row


# ------------------------------------------------------------------ Word
def render_docx(out_path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Cm

    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)

    for block in BLOCKS:
        kind = block[0]
        if kind == 'title':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(block[1]); r.bold = True; r.font.size = Pt(26); r.font.color.rgb = ACCENT
        elif kind == 'subtitle':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(block[1]); r.font.size = Pt(14); r.font.color.rgb = GREY
        elif kind == 'meta':
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(block[1]); r.font.size = Pt(9); r.font.color.rgb = GREY
        elif kind == 'h1':
            doc.add_heading(block[1], level=1)
        elif kind == 'h2':
            doc.add_heading(block[1], level=2)
        elif kind == 'h3':
            doc.add_heading(block[1], level=3)
        elif kind == 'p':
            doc.add_paragraph(block[1])
        elif kind == 'bullets':
            for item in block[1]:
                p = doc.add_paragraph(style='List Bullet')
                if isinstance(item, tuple):
                    r = p.add_run(item[0]); r.bold = True; p.add_run(item[1])
                else:
                    p.add_run(item)
        elif kind == 'callout':
            p = doc.add_paragraph(style='Intense Quote')
            p.add_run(block[1])
        elif kind == 'table':
            headers, rows, widths = block[1], block[2], block[3]
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = 'Light Grid Accent 1'
            for i, h in enumerate(headers):
                t.rows[0].cells[i].text = h
            style_header_row(t.rows[0])
            for row in rows:
                cells = t.add_row().cells
                for i, val in enumerate(row):
                    cells[i].text = str(val)
            for r in t.rows:
                for i, w in enumerate(widths):
                    r.cells[i].width = Cm(w)
        elif kind == 'code':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(block[1])
            run.font.name = 'Consolas'
            run.font.size = Pt(8.5)
        elif kind == 'pagebreak':
            doc.add_page_break()

    doc.save(out_path)
    print(f"Word généré : {out_path}")


# ------------------------------------------------------------------ PDF
def render_pdf(out_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib.enums import TA_CENTER
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Table, TableStyle, ListFlowable, ListItem, PageBreak,
    )

    accent = colors.HexColor('#1F4E79')
    grey = colors.HexColor('#595959')
    light = colors.HexColor('#DCE6F1')
    codebg = colors.HexColor('#F2F2F2')

    ss = getSampleStyleSheet()
    ss.add(ParagraphStyle('DocTitle', parent=ss['Title'], fontSize=26, textColor=accent, spaceAfter=4))
    ss.add(ParagraphStyle('DocSub', parent=ss['Normal'], fontSize=14, textColor=grey, alignment=TA_CENTER))
    ss.add(ParagraphStyle('DocMeta', parent=ss['Normal'], fontSize=9, textColor=grey, alignment=TA_CENTER))
    ss.add(ParagraphStyle('H1', parent=ss['Heading1'], fontSize=15, textColor=accent, spaceBefore=14, spaceAfter=6))
    ss.add(ParagraphStyle('H2', parent=ss['Heading2'], fontSize=12, textColor=accent, spaceBefore=10, spaceAfter=4))
    ss.add(ParagraphStyle('H3', parent=ss['Heading3'], fontSize=11, textColor=accent, spaceBefore=8, spaceAfter=3))
    ss.add(ParagraphStyle('Body', parent=ss['Normal'], fontSize=10.5, leading=15, spaceAfter=6))
    ss.add(ParagraphStyle('CalloutBox', parent=ss['Normal'], fontSize=10.5, leading=15,
                          textColor=accent, leftIndent=10, spaceBefore=6, spaceAfter=8,
                          borderPadding=6, backColor=light))
    ss.add(ParagraphStyle('CodeBlk', parent=ss['Normal'], fontName='Courier', fontSize=8.5,
                          leading=11, backColor=codebg, leftIndent=8, spaceBefore=4, spaceAfter=8,
                          borderPadding=6))

    def make_table(headers, rows, widths):
        data = [headers] + [[str(c) for c in r] for r in rows]
        t = Table(data, colWidths=[w * cm for w in widths], repeatRows=1, hAlign='CENTER')
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), accent),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 8.5),
            ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#9DB7D4')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, light]),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        return t

    story = []
    for block in BLOCKS:
        kind = block[0]
        if kind == 'title':
            story.append(Paragraph(block[1], ss['DocTitle']))
        elif kind == 'subtitle':
            story.append(Paragraph(block[1], ss['DocSub']))
        elif kind == 'meta':
            story.append(Paragraph(block[1], ss['DocMeta']))
        elif kind in ('h1', 'h2', 'h3'):
            story.append(Paragraph(block[1], ss[kind.upper()]))
        elif kind == 'p':
            story.append(Paragraph(block[1], ss['Body']))
        elif kind == 'bullets':
            lis = []
            for item in block[1]:
                if isinstance(item, tuple):
                    lead, rest = item
                    # escape lead/rest for XML, bold the lead via <b>
                    lead_e = lead.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    rest_e = rest.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
                    text = f'<b>{lead_e}</b>{rest_e}'
                else:
                    text = str(item)
                lis.append(ListItem(Paragraph(text, ss['Body']), leftIndent=12))
            story.append(ListFlowable(lis, bulletType='bullet', start='•', leftIndent=10))
        elif kind == 'callout':
            story.append(Paragraph(block[1], ss['CalloutBox']))
        elif kind == 'table':
            story.append(make_table(block[1], block[2], block[3]))
        elif kind == 'code':
            code = block[1].replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')
            code = code.replace('\n', '<br/>')
            story.append(Paragraph(code, ss['CodeBlk']))
        elif kind == 'pagebreak':
            story.append(PageBreak())

    def _decor(canvas, doc):
        canvas.saveState()
        canvas.setFont('Helvetica', 8)
        canvas.setFillColor(grey)
        canvas.drawString(2 * cm, 1.2 * cm, 'AirLLM — Compte rendu de sécurité')
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f'Page {doc.page}')
        canvas.setStrokeColor(light)
        canvas.line(2 * cm, 1.5 * cm, A4[0] - 2 * cm, 1.5 * cm)
        canvas.restoreState()

    doc = SimpleDocTemplate(
        out_path, pagesize=A4,
        leftMargin=1.8 * cm, rightMargin=1.8 * cm, topMargin=1.8 * cm, bottomMargin=2.2 * cm,
        title='AirLLM — Compte rendu de sécurité', author='Audit sécurité AirLLM')
    doc.build(story, onFirstPage=_decor, onLaterPages=_decor)
    print(f"PDF généré : {out_path}")


def main(base):
    render_docx(base + '.docx')
    render_pdf(base + '.pdf')


if __name__ == '__main__':
    base = sys.argv[1] if len(sys.argv) > 1 else 'AirLLM_Compte_Rendu_Securite_Detaille'
    main(base)

